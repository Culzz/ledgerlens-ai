"""Tests for the workbook export.

The file is the one artefact a reader takes away, so these check the things that
would make it a wrong artefact rather than an ugly one: that a figure the report does
not state is labelled as such, that the exact decimal is not rounded on its way into
a spreadsheet cell, and that the sheet layout does not quietly imply a step is three
steps because it has three arguments.
"""

from __future__ import annotations

import zipfile
from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO

from app.services.reasoning import SOURCED, UNSOURCED, AuditStep, Operand
from app.services.reporting import AnswerReport, EvidenceRow, filename_for, workbook_bytes
from openpyxl import load_workbook

RATIO = "1.082969432314410480349344978"


def steps() -> tuple[AuditStep, ...]:
    # Operand text as the audit really writes it: a figure from the report carries
    # the row it was found in, and a figure that is nowhere carries the mark saying
    # so. The sheet copies those, so a fixture that invented a tidier spelling would
    # be testing a layout nobody will ever see.
    return (
        AuditStep(
            index=0,
            operation="divide",
            operands=(
                Operand(
                    text="24800 [table_2]",
                    role=SOURCED,
                    value=Decimal(24800),
                    labels=("table_2",),
                ),
                Operand(text="22900 (unverified)", role=UNSOURCED, value=Decimal(22900)),
            ),
            result=Decimal(RATIO),
        ),
    )


def report(**overrides) -> AnswerReport:
    defaults: dict = {
        "question": "what is the ratio of total flight attendants to passenger service personnel",
        "document_id": "aal-2013-page-15-pdf",
        "filename": "AAL/2013/page_15.pdf",
        "answer": RATIO,
        "program": "divide(24800, 22900)",
        "failure": None,
        "verdict": "partly_grounded",
        "unsourced": ("22900",),
        "sourced_count": 1,
        "literal_count": 2,
        "report_unit": None,
        "steps": steps(),
        "evidence": (
            EvidenceRow(label="table_0", kind="table", cited=False, text="american us airways total"),
            EvidenceRow(label="table_2", kind="table", cited=True, text="flight attendants 24800"),
        ),
        "provider": "mistral",
        "model": "ministral-14b-2512",
        "generated_at": datetime(2026, 10, 19, 9, 30, tzinfo=UTC),
    }
    return AnswerReport(**{**defaults, **overrides})


def sheet(data: bytes, name: str) -> list[list[str | None]]:
    book = load_workbook(BytesIO(data))
    return [[cell.value for cell in row] for row in book[name].iter_rows()]


def answer_field(data: bytes, field: str) -> str | None:
    for row in sheet(data, "Answer"):
        if row and row[0] == field:
            return row[1]
    raise AssertionError(f"no {field!r} row in the Answer sheet")


class TestWorkbook:
    def test_it_is_a_workbook_with_the_three_sheets(self):
        book = load_workbook(BytesIO(workbook_bytes(report())))
        assert book.sheetnames == ["Answer", "Trail", "Evidence"]

    def test_the_answer_sheet_carries_the_question_the_figure_and_the_program(self):
        data = workbook_bytes(report())

        assert answer_field(data, "Question") == report().question
        assert answer_field(data, "Answer") == RATIO
        assert answer_field(data, "Program") == "divide(24800, 22900)"

    def test_the_verdict_is_written_in_words_rather_than_as_an_enum(self):
        assert answer_field(workbook_bytes(report()), "Check") == "partly grounded"

    def test_a_figure_the_report_does_not_state_is_named_as_one(self):
        # The reason the audit exists. A workbook that printed the bare figure and
        # said nothing about where it came from would be the failing this whole
        # project is built against.
        assert answer_field(workbook_bytes(report()), "Not in the report") == "22900"

    def test_a_grounded_answer_has_no_such_row(self):
        data = workbook_bytes(report(verdict="grounded", unsourced=(), literal_count=1))

        assert all(row[0] != "Not in the report" for row in sheet(data, "Answer") if row)

    def test_the_exact_decimal_is_not_rounded_into_a_spreadsheet_number(self):
        written = answer_field(workbook_bytes(report()), "Answer")

        assert written == RATIO
        assert isinstance(written, str)

    def test_the_digits_are_in_the_file_itself(self):
        # Not only the cell this library hands back. If anything between the engine
        # and the sheet had parsed the figure as a float, the file would hold 15
        # digits and the cell would still read back looking right.
        archive = zipfile.ZipFile(BytesIO(workbook_bytes(report())))
        blob = b"".join(archive.read(name) for name in archive.namelist())

        assert RATIO.encode() in blob

    def test_a_program_that_never_ran_says_so_instead_of_leaving_a_gap(self):
        data = workbook_bytes(
            report(answer=None, program=None, failure="divide: division by zero", steps=())
        )

        assert answer_field(data, "Answer") == "no answer"
        assert answer_field(data, "Why there is no answer") == "divide: division by zero"


class TestTrailSheet:
    def test_every_figure_gets_a_row_and_says_where_it_came_from(self):
        rows = sheet(workbook_bytes(report()), "Trail")

        assert rows[0] == [
            "Step",
            "Operation",
            "Result",
            "Step as written",
            "Operand",
            "Where the figure came from",
            "Rows",
        ]
        assert [row[4] for row in rows[1:]] == ["24800 [table_2]", "22900 (unverified)"]
        assert [row[5] for row in rows[1:]] == ["from the report", "not in the report"]

    def test_the_steps_own_columns_are_written_once_for_a_step_with_two_figures(self):
        # Two operands are one step. Repeating the result beside each of them would
        # read as two steps that happened to agree.
        rows = sheet(workbook_bytes(report()), "Trail")

        assert rows[1][:4] == [
            0,
            "divide",
            RATIO,
            f"#0 = divide(24800 [table_2], 22900 (unverified)) = {RATIO}",
        ]
        assert rows[2][:4] == [None, None, None, None]

    def test_the_labels_a_figure_was_found_in_are_kept(self):
        rows = sheet(workbook_bytes(report()), "Trail")

        assert rows[1][6] == "table_2"
        assert rows[2][6] is None

    def test_a_step_with_no_figures_still_appears(self):
        data = workbook_bytes(
            report(steps=(AuditStep(index=0, operation="", operands=(), result="yes"),))
        )

        assert sheet(data, "Trail")[1] == [0, None, "yes", "#0 = yes", None, None, None]


class TestEvidenceSheet:
    def test_a_row_that_was_offered_and_not_used_says_so(self):
        rows = sheet(workbook_bytes(report()), "Evidence")

        assert [row[2] for row in rows[1:]] == ["offered", "used"]
        assert rows[1][0] == "table_0"
        assert rows[2][3] == "flight attendants 24800"

    def test_an_empty_retrieval_says_so_rather_than_showing_a_bare_header(self):
        rows = sheet(workbook_bytes(report(evidence=(), steps=())), "Evidence")

        assert rows[1][3] == "nothing was retrieved for this question"


class TestFilename:
    def test_it_is_built_from_the_document_and_the_day(self):
        assert filename_for("aal-2013-page-15-pdf", datetime(2026, 10, 19, tzinfo=UTC)) == (
            "ledgerlens-aal-2013-page-15-pdf-2026-10-19.xlsx"
        )

    def test_a_document_id_with_a_slash_in_it_cannot_escape_the_filename(self):
        assert filename_for("../../etc/passwd", datetime(2026, 10, 19, tzinfo=UTC)) == (
            "ledgerlens-etc-passwd-2026-10-19.xlsx"
        )

    def test_an_id_with_nothing_usable_in_it_falls_back(self):
        assert filename_for("///", datetime(2026, 10, 19, tzinfo=UTC)) == (
            "ledgerlens-answer-2026-10-19.xlsx"
        )
